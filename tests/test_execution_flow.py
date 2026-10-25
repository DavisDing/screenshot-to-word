"""Qt flow checks with real Excel/Word services; no screen capture or native GUI acceptance."""
from __future__ import annotations

import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from PySide6.QtWidgets import QApplication, QDialog

from qt_app.services.logger import AppLogger
from qt_app.ui.main_window import MainWindow
from qt_app.ui.excel_mapping_dialog import ExcelMappingDialog


class ExecutionFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)

    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.base = Path(self.temporary.name)
        self.path = self.base / "two_sheets.xlsx"
        workbook = Workbook()
        workbook.active.title = "第一张"
        workbook.active.append(["测试名称", "验证点", "测试结果"])
        workbook.active.append(["登录", "登录成功", None])
        second = workbook.create_sheet("第二张")
        second.append(["Status", "Case Name", "Check Point"])
        second.append([None, "支付", "支付成功"])
        workbook.active = 1
        workbook.save(self.path)
        workbook.close()
        self.window = MainWindow(self.base, AppLogger(self.base / "logs/output.log"))
        self.addCleanup(self.cleanup_window)

    def cleanup_window(self):
        if self.window.control_panel is not None:
            self.window.control_panel.reject()
        if self.window.excel_service is not None:
            self.window.excel_service.close()
        self.window.close()
        self.app.processEvents()
        self.app.setQuitOnLastWindowClosed(False)

    def test_mapping_execution_and_word_excel_persistence_cross_sheet_boundary(self):
        dialogs = []

        def confirm(dialog):
            dialogs.append(dialog.preview.sheet_name)
            return QDialog.DialogCode.Accepted

        with patch("qt_app.ui.main_window.QFileDialog.getOpenFileName", return_value=(str(self.path), "")), \
             patch.object(ExcelMappingDialog, "exec", confirm):
            self.window.select_and_run()
        panel = self.window.control_panel
        self.assertEqual(dialogs, ["第一张", "第二张"])
        self.assertEqual([(item.sheet_name, item.row_index) for item in panel.items], [("第一张", 2), ("第二张", 2)])
        image = self.base / "annotated.png"
        Image.new("RGB", (80, 40), "blue").save(image)
        with patch("qt_app.ui.control_window.QMessageBox.information") as finished:
            for sheet in dialogs:
                item = panel._current()
                self.assertEqual(item.sheet_name, sheet)
                self.assertIn(sheet, panel.windowTitle())
                panel.last_image = image
                self.assertTrue(panel._save_current_result(item, image))
                panel.finish_current()
            self.assertEqual(finished.call_count, 1)
        self.assertIsNone(self.window.control_panel)
        saved = load_workbook(self.path)
        try:
            self.assertEqual(saved["第一张"]["C2"].value, "已执行")
            self.assertEqual(saved["第二张"]["A2"].value, "已执行")
            self.assertEqual(saved["第二张"]["B2"].value, "支付")
        finally:
            saved.close()
        for case in ("登录", "支付"):
            report = Document(self.base / "word_output" / f"{case}.docx")
            self.assertEqual(report.paragraphs[0].text, case)
            self.assertEqual(len(report.inline_shapes), 1)

    def test_same_case_name_across_sheets_has_separate_word_reports(self):
        workbook = load_workbook(self.path)
        workbook["第一张"]["A2"] = "共同用例"
        workbook["第二张"]["B2"] = "共同用例"
        workbook.save(self.path)
        workbook.close()
        with patch("qt_app.ui.main_window.QFileDialog.getOpenFileName", return_value=(str(self.path), "")), \
             patch.object(ExcelMappingDialog, "exec", return_value=QDialog.DialogCode.Accepted):
            self.window.select_and_run()
        panel = self.window.control_panel
        image = self.base / "annotated.png"
        Image.new("RGB", (80, 40), "blue").save(image)
        panel.last_image = image
        first = panel._current()
        self.assertTrue(panel._save_current_result(first, image))
        panel.finish_current()
        panel.last_image = image
        second = panel._current()
        self.assertTrue(panel._save_current_result(second, image))
        panel.finish_current()
        reports = list((self.base / "word_output").glob("*.docx"))
        self.assertEqual(len(reports), 2)
        self.assertTrue(all(len(Document(path).inline_shapes) == 1 for path in reports))

    def test_cancel_second_mapping_creates_no_execution_or_partial_writes(self):
        original = self.path.read_bytes()
        with patch("qt_app.ui.main_window.QFileDialog.getOpenFileName", return_value=(str(self.path), "")), \
             patch.object(ExcelMappingDialog, "exec", side_effect=[QDialog.DialogCode.Accepted, QDialog.DialogCode.Rejected]):
            self.window.select_and_run()
        self.assertIsNone(self.window.control_panel)
        self.assertIsNone(self.window.excel_service.workbook)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list((self.base / "word_output").glob("*.docx")), [])

    def test_skip_and_end_leave_other_sheets_pending(self):
        original = self.path.read_bytes()
        with patch("qt_app.ui.main_window.QFileDialog.getOpenFileName", return_value=(str(self.path), "")), \
             patch.object(ExcelMappingDialog, "exec", return_value=QDialog.DialogCode.Accepted):
            self.window.select_and_run()
        panel = self.window.control_panel
        panel.skip_current()
        self.assertEqual(panel._current().sheet_name, "第二张")
        panel.reject()
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list((self.base / "word_output").glob("*.docx")), [])


if __name__ == "__main__":
    unittest.main()
