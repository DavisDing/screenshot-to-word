from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from qt_app.services.excel_service import ExcelService
from qt_app.services.logger import AppLogger
from qt_app.services.word_service import WordService


def _make_image(path: Path) -> None:
    img = Image.new("RGB", (400, 240), "white")
    ImageDraw.Draw(img).rectangle((20, 20, 180, 120), outline="red", width=4)
    img.save(path)


def test_excel_mapping_and_word_generation() -> None:
    with TemporaryDirectory() as td:
        base = Path(td)
        logger = AppLogger(base / "logs" / "output.log")

        xlsx = base / "cases.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["Case Name", "Check Point", "Step", "Desc", "Expected", "Status"])
        ws.append(["登录功能", "登录成功", "打开页面", "点击登录按钮", "显示首页", ""])
        ws.append(["已完成用例", "不应返回", "", "", "", "已执行"])
        wb.save(xlsx)

        excel = ExcelService(logger)
        preview = excel.preview(xlsx)
        assert preview.suggested_mapping["测试名称"] == "Case Name"
        assert preview.suggested_mapping["验证点"] == "Check Point"
        assert preview.suggested_mapping["测试结果"] == "Status"

        items = excel.load(xlsx, preview.suggested_mapping)
        assert len(items) == 1
        assert items[0].case_name == "登录功能"
        assert items[0].step_name == "打开页面"
        excel.mark_executed([items[0].row_index])
        excel.close()

        wb2 = load_workbook(xlsx)
        assert wb2.active["F2"].value == "已执行"
        wb2.close()

        image_path = base / "shot.png"
        _make_image(image_path)

        word = WordService(base / "word_output", logger)
        doc_path = word.add_image("登录/功能:*?", "登录成功", image_path, "步骤名称：打开页面")
        assert doc_path.exists()
        assert doc_path.name == "登录_功能.docx"


def test_merged_case_rows_execute_step_by_step() -> None:
    with TemporaryDirectory() as td:
        base = Path(td)
        logger = AppLogger(base / "logs" / "output.log")
        xlsx = base / "merged_steps.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.append(["测试名称", "验证点", "步骤名称", "步骤描述", "预期结果", "测试结果"])
        ws.merge_cells("A2:A4")
        ws.merge_cells("B2:B4")
        ws["A2"] = "合并单元格用例"
        ws["B2"] = "合并验证点"
        ws["C2"] = "前置条件"
        ws["D2"] = "数据运维运行正常"
        ws["E2"] = "运行正常"
        ws["F2"] = "已执行"
        ws["C3"] = "1"
        ws["D3"] = "配置规则"
        ws["E3"] = "配置成功"
        ws["C4"] = "2"
        ws["D4"] = "触发告警"
        ws["E4"] = "收到告警"
        wb.save(xlsx)

        excel = ExcelService(logger)
        preview = excel.preview(xlsx)
        items = excel.load(xlsx, preview.suggested_mapping)

        assert [item.row_index for item in items] == [3, 4]
        assert [item.step_name for item in items] == ["1", "2"]
        assert all(item.case_name == "合并单元格用例" for item in items)
        assert all(item.checkpoint == "合并验证点" for item in items)

        excel.mark_executed([item.row_index for item in items])
        excel.close()

        wb2 = load_workbook(xlsx)
        ws2 = wb2.active
        assert ws2["F3"].value == "已执行"
        assert ws2["F4"].value == "已执行"
        wb2.close()


def test_word_checkpoint_is_not_duplicated() -> None:
    with TemporaryDirectory() as td:
        base = Path(td)
        logger = AppLogger(base / "logs" / "output.log")
        image_path = base / "shot.png"
        _make_image(image_path)

        word = WordService(base / "word_output", logger)
        doc_path = word.add_image("同一用例", "同一验证点", image_path, "步骤名称：1")
        word.add_image("同一用例", "同一验证点", image_path, "步骤名称：2")

        doc = Document(str(doc_path))
        checkpoint_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip() == "验证点：同一验证点"]
        assert len(checkpoint_lines) == 1


if __name__ == "__main__":
    test_excel_mapping_and_word_generation()
    test_merged_case_rows_execute_step_by_step()
    test_word_checkpoint_is_not_duplicated()
    print("smoke tests passed")
