from __future__ import annotations

import hashlib
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest
import unicodedata

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from qt_app.services.logger import AppLogger
from qt_app.services.screenshot_service import ScreenshotService
from qt_app.services.word_service import WordService


class WordReportRegressionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.base = Path(self.temporary.name)
        self.output = self.base / "reports"
        self.logger = AppLogger(self.base / "output.log")
        self.addCleanup(self.logger.close)
        self.image = self.base / "shot.png"
        Image.new("RGB", (40, 30), "white").save(self.image)

    def add_image(self, case_name: str, checkpoint: str = "验证点1") -> Path:
        # Every call constructs a service, as when the application is restarted.
        return WordService(self.output, self.logger).add_image(case_name, checkpoint, self.image)

    def add_scoped_image(self, case_name: str, checkpoint: str, *, sheet_name: str, operation_id: str = "") -> Path:
        return WordService(self.output, self.logger).add_image(
            case_name, checkpoint, self.image, sheet_name=sheet_name, operation_id=operation_id,
        )

    def legacy_report(self, filename: str, heading: str | None, identifier: str = "") -> Path:
        self.output.mkdir(exist_ok=True)
        path = self.output / filename
        doc = Document()
        if heading is not None:
            doc.add_heading(heading, level=1)
        doc.add_paragraph("已有内容")
        doc.core_properties.identifier = identifier
        doc.save(path)
        return path

    def test_sanitized_names_produce_separate_reports(self) -> None:
        names = ["登录/功能", "登录:功能", "登录?功能"]
        paths = [self.add_image(name, f"验证点{index}") for index, name in enumerate(names, 1)]

        self.assertEqual(paths[0].name, "登录_功能.docx")
        self.assertEqual(len(set(paths)), 3)
        for index, (name, path) in enumerate(zip(names, paths), 1):
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].text, name)
            self.assertEqual([p.text for p in doc.paragraphs if p.text.startswith("验证点：")],
                             [f"验证点：验证点{index}"])
            self.assertEqual(len(doc.inline_shapes), 1)

    def test_truncated_names_produce_separate_reports(self) -> None:
        names = ["A" * 80 + "一", "A" * 80 + "二"]
        paths = [self.add_image(name) for name in names]

        self.assertNotEqual(paths[0], paths[1])
        self.assertEqual(paths[0].name, "A" * 80 + ".docx")
        for name, path in zip(names, paths):
            self.assertLessEqual(len(path.stem), 80)
            self.assertEqual(self.add_image(name), path)
            self.assertEqual(len(Document(path).inline_shapes), 2)

    def test_names_that_differ_only_in_trimmed_characters_stay_separate(self) -> None:
        paths = [self.add_image(name) for name in ["登录", "__登录__", " .登录. "]]

        self.assertEqual(len(set(paths)), 3)
        self.assertEqual(paths[0].name, "登录.docx")

    def test_case_and_unicode_equivalent_filenames_stay_separate(self) -> None:
        names = ["Login", "login", "café", "cafe\u0301"]
        paths = [self.add_image(name) for name in names]
        portable_names = [unicodedata.normalize("NFC", path.name).casefold() for path in paths]

        self.assertEqual(len(set(portable_names)), len(names))
        for name, path in zip(names, paths):
            self.assertEqual(self.add_image(name), path)
            self.assertEqual(Document(path).paragraphs[0].text, name)

    def test_recreated_service_appends_to_own_report_and_deduplicates_checkpoint(self) -> None:
        first = self.add_image("登录/功能", "同一验证点")
        second = self.add_image("登录:功能", "同一验证点")

        self.assertEqual(self.add_image("登录:功能", "同一验证点"), second)
        self.assertEqual(self.add_image("登录/功能", "同一验证点"), first)
        for path in (first, second):
            doc = Document(path)
            self.assertEqual(len(doc.inline_shapes), 2)
            self.assertEqual([p.text for p in doc.paragraphs].count("验证点：同一验证点"), 1)

    def test_matching_legacy_report_is_reused_and_gets_persistent_identity(self) -> None:
        path = self.legacy_report("登录_功能.docx", "登录/功能")

        self.assertEqual(self.add_image("登录/功能"), path)
        doc = Document(path)
        self.assertIn("已有内容", [p.text for p in doc.paragraphs])
        self.assertTrue(doc.core_properties.identifier)
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_foreign_or_unidentified_reports_are_preserved(self) -> None:
        for heading, identifier in [("其他用例", ""), (None, ""), ("登录/功能", "foreign-owner")]:
            with self.subTest(heading=heading, identifier=identifier):
                # Separate each scenario without deleting an existing report.
                self.output = self.base / f"reports-{len(list(self.base.iterdir()))}"
                path = self.legacy_report("登录_功能.docx", heading, identifier)
                original = path.read_bytes()

                result = self.add_image("登录/功能")

                self.assertNotEqual(result, path)
                self.assertEqual(path.read_bytes(), original)
                self.assertEqual(Document(result).paragraphs[0].text, "登录/功能")

    def test_unreadable_existing_file_is_preserved(self) -> None:
        self.output.mkdir()
        path = self.output / "登录_功能.docx"
        path.write_bytes(b"not a Word document")

        result = self.add_image("登录/功能")

        self.assertNotEqual(result, path)
        self.assertEqual(path.read_bytes(), b"not a Word document")

    def test_occupied_digest_candidate_is_preserved_and_own_report_is_reused_after_gaps(self) -> None:
        name = "登录:功能"
        digest = hashlib.sha256(name.encode("utf-8")).hexdigest()[:12]
        base_path = self.legacy_report("登录_功能.docx", "登录/功能")
        digest_path = self.legacy_report(f"登录_功能-{digest}.docx", "其他用例")
        base_bytes, digest_bytes = base_path.read_bytes(), digest_path.read_bytes()

        path = self.add_image(name)

        self.assertEqual(path.name, f"登录_功能-{digest}-2.docx")
        self.assertEqual(base_path.read_bytes(), base_bytes)
        self.assertEqual(digest_path.read_bytes(), digest_bytes)
        base_path.unlink()
        digest_path.unlink()
        self.assertEqual(self.add_image(name), path)
        self.assertEqual(len(Document(path).inline_shapes), 2)

    def test_original_newlines_are_distinguished_across_restarts(self) -> None:
        names = ["登录\r\n功能", "登录\n功能"]
        paths = [self.add_image(name) for name in names]

        self.assertNotEqual(paths[0], paths[1])
        for name, path in zip(names, paths):
            self.assertEqual(self.add_image(name), path)
            self.assertEqual(len(Document(path).inline_shapes), 2)
        self.assertNotEqual(Document(paths[0]).core_properties.identifier,
                            Document(paths[1]).core_properties.identifier)

    def test_same_case_on_different_sheets_uses_separate_reports(self) -> None:
        first = self.add_scoped_image("登录/成功", "成功", sheet_name="Web")
        second = self.add_scoped_image("登录/成功", "成功", sheet_name="API")
        self.assertNotEqual(first, second)
        self.assertEqual(len(Document(first).inline_shapes), 1)
        self.assertEqual(len(Document(second).inline_shapes), 1)
        ids = {Document(first).core_properties.identifier, Document(second).core_properties.identifier}
        self.assertEqual(len(ids), 2)

    def test_operation_id_makes_retry_idempotent_after_excel_failure(self) -> None:
        path = self.add_scoped_image("登录", "成功", sheet_name="Web", operation_id="stable-operation")
        retry = self.add_scoped_image("登录", "成功", sheet_name="Web", operation_id="stable-operation")
        doc = Document(path)
        self.assertEqual(path, retry)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual(doc.inline_shapes[0]._inline.docPr.get("descr"), "screenshot-to-word-entry:stable-operation")

    def test_repeated_checkpoint_images_stay_before_next_checkpoint(self) -> None:
        service = WordService(self.output, self.logger)
        service.add_image("登录", "验证点1", self.image, "步骤1")
        service.add_image("登录", "验证点2", self.image, "步骤2")
        path = service.add_image("登录", "验证点1", self.image, "步骤3")
        doc = Document(path)
        text = [p.text for p in doc.paragraphs]
        p1 = text.index("验证点：验证点1")
        p2 = text.index("验证点：验证点2")
        self.assertLess(p1, text.index("步骤3"))
        self.assertLess(text.index("步骤3"), p2)
        self.assertEqual(len(doc.inline_shapes), 3)

    def test_inserted_image_width_fits_a4_default_text_area(self) -> None:
        path = self.add_image("宽度")
        width_emu = Document(path).inline_shapes[0].width
        self.assertLessEqual(width_emu, int(5.6 * 914400))

    def test_temp_cleanup_removes_only_stale_managed_pngs(self) -> None:
        temp_dir = self.base / "Temp"
        temp_dir.mkdir()
        stale = temp_dir / "stale.png"
        recent = temp_dir / "recent.png"
        nested_dir = temp_dir / "nested"
        nested_dir.mkdir()
        nested = nested_dir / "keep.png"
        for path in (stale, recent, nested):
            path.write_bytes(b"image")
        import os, time
        os.utime(stale, (time.time() - 8 * 86400, time.time() - 8 * 86400))

        ScreenshotService(temp_dir, self.logger)
        self.assertFalse(stale.exists())
        self.assertTrue(recent.exists())
        self.assertTrue(nested.exists())


if __name__ == "__main__":
    unittest.main()
