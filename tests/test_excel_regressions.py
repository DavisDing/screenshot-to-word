from __future__ import annotations

from pathlib import Path
from shutil import copyfileobj
from tempfile import TemporaryDirectory
import sys
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import PatternFill
from lxml import etree
from PIL import Image

from qt_app.services.excel_service import ExcelService


HEADERS = ["测试名称", "验证点", "步骤名称", "步骤描述", "预期结果", "测试结果"]
MAPPING = {header: header for header in HEADERS}
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
VBA_CONTENT_TYPE = "application/vnd.ms-office.vbaProject"
MACRO_WORKBOOK_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
VBA_REL_TYPE = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
VBA_BYTES = b"synthetic-vba-project\x00\xff\x01preserve-verbatim"


class SilentLogger:
    def log(self, message: str) -> None:
        pass


def make_workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    workbook.active.append(HEADERS)
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


def add_macro_package(path: Path) -> None:
    """Add an inert VBA part and the package declarations used by macro workbooks."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for entry in content_types:
        if entry.get("PartName") == "/xl/workbook.xml":
            entry.set("ContentType", MACRO_WORKBOOK_TYPE)
    ET.SubElement(content_types, f"{{{CONTENT_NS}}}Default", {
        "Extension": "bin", "ContentType": VBA_CONTENT_TYPE,
    })
    parts["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8")

    relationships = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    ET.SubElement(relationships, f"{{{REL_NS}}}Relationship", {
        "Id": "rIdVba", "Type": VBA_REL_TYPE, "Target": "vbaProject.bin",
    })
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(relationships, encoding="utf-8")
    parts["xl/vbaProject.bin"] = VBA_BYTES
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


class ExcelRegressionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.base = Path(self.temporary_directory.name)
        self.excel = ExcelService(SilentLogger())
        self.addCleanup(self.excel.close)

    def assert_macro_package_preserved(self, path: Path) -> None:
        with ZipFile(path) as archive:
            self.assertIn("xl/vbaProject.bin", archive.namelist())
            self.assertEqual(archive.read("xl/vbaProject.bin"), VBA_BYTES)
            content_types = ET.fromstring(archive.read("[Content_Types].xml"))
            self.assertTrue(any(
                entry.get("PartName") == "/xl/workbook.xml"
                and entry.get("ContentType") == MACRO_WORKBOOK_TYPE
                for entry in content_types
            ))
            self.assertTrue(any(
                (entry.get("Extension") == "bin" or entry.get("PartName") == "/xl/vbaProject.bin")
                and entry.get("ContentType") == VBA_CONTENT_TYPE
                for entry in content_types
            ))
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            self.assertTrue(any(
                entry.get("Type") == VBA_REL_TYPE
                and entry.get("Target") in ("vbaProject.bin", "/xl/vbaProject.bin")
                for entry in relationships
            ))

    def test_macro_package_survives_repeated_status_saves_and_reload(self) -> None:
        for suffix in (".xlsm", ".XLSM", ".XlSm"):
            with self.subTest(suffix=suffix):
                path = self.base / f"cases{suffix}"
                make_workbook(path, [
                    ["登录", "成功", "步骤1", "", "", ""],
                    ["", "", "步骤2", "", "", ""],
                ])
                add_macro_package(path)
                original_bytes = path.read_bytes()
                preview = self.excel.preview(path)
                self.assertEqual(path.read_bytes(), original_bytes)
                items = self.excel.load(path, preview.suggested_mapping)
                self.assertEqual([item.row_index for item in items], [2, 3])
                self.assertEqual(path.read_bytes(), original_bytes)

                self.excel.mark_executed([2])
                self.assert_macro_package_preserved(path)
                self.excel.mark_executed([3])
                self.assert_macro_package_preserved(path)
                saved_workbook = self.excel.workbook
                self.assertEqual(saved_workbook.active["F2"].value, "已执行")
                self.assertEqual(saved_workbook.active["F3"].value, "已执行")

                self.assertEqual(self.excel.load(path, MAPPING), [])
                self.assertIsNone(saved_workbook.vba_archive.fp)
                reloaded_workbook = self.excel.workbook
                self.excel.close()
                self.excel.close()
                self.assertIsNone(reloaded_workbook.vba_archive.fp)
                self.assert_macro_package_preserved(path)

    def test_macro_preview_releases_archive_even_when_preview_fails(self) -> None:
        path = self.base / "preview.xlsm"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        add_macro_package(path)
        original_bytes = path.read_bytes()
        loaded_workbooks = []

        def capture_workbook(*args, **kwargs):
            workbook = load_workbook(*args, **kwargs)
            loaded_workbooks.append(workbook)
            return workbook

        with patch("qt_app.services.excel_service.load_workbook", side_effect=capture_workbook):
            self.excel.preview(path)
            with patch.object(self.excel, "_detect_header_row", side_effect=ValueError("invalid preview")):
                with self.assertRaisesRegex(ValueError, "invalid preview"):
                    self.excel.preview(path)
        self.assertEqual(len(loaded_workbooks), 2)
        for workbook in loaded_workbooks:
            self.assertIsNotNone(workbook.vba_archive)
            self.assertIsNone(workbook.vba_archive.fp)
        self.assertEqual(path.read_bytes(), original_bytes)

    def test_invalid_macro_mapping_releases_archive(self) -> None:
        path = self.base / "mapping.xlsm"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        add_macro_package(path)
        loaded_workbooks = []

        def capture_workbook(*args, **kwargs):
            workbook = load_workbook(*args, **kwargs)
            loaded_workbooks.append(workbook)
            return workbook

        with patch("qt_app.services.excel_service.load_workbook", side_effect=capture_workbook):
            with self.assertRaisesRegex(ValueError, "Excel 缺少必要列映射"):
                self.excel.load(path, {})
        self.assertIsNotNone(loaded_workbooks[0].vba_archive)
        self.assertIsNone(loaded_workbooks[0].vba_archive.fp)
        self.assertIsNone(self.excel.workbook)

    def test_xlsx_status_save_remains_an_ordinary_workbook(self) -> None:
        path = self.base / "ordinary.xlsx"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        original_bytes = path.read_bytes()
        preview = self.excel.preview(path)
        self.assertEqual(path.read_bytes(), original_bytes)
        items = self.excel.load(path, preview.suggested_mapping)
        self.assertIsNone(self.excel.workbook.vba_archive)
        self.excel.mark_executed([item.row_index for item in items])
        self.excel.close()
        with ZipFile(path) as archive:
            self.assertNotIn("xl/vbaProject.bin", archive.namelist())
        workbook = load_workbook(path)
        try:
            self.assertEqual(workbook.active["F2"].value, "已执行")
        finally:
            workbook.close()

    def test_repeated_saves_preserve_images_on_all_sheets_and_macros(self) -> None:
        image_path = self.base / "existing.png"
        Image.new("RGB", (20, 16), "blue").save(image_path)
        for suffix in (".xlsx", ".xlsm"):
            with self.subTest(suffix=suffix):
                path = self.base / f"images{suffix}"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(HEADERS)
                sheet.append(["登录", "成功", "步骤1", "", "", ""])
                sheet.append(["", "", "步骤2", "", "", ""])
                sheet.add_image(WorksheetImage(str(image_path)), "H2")
                other = workbook.create_sheet("其他工作表")
                other["A1"] = "保留内容"
                other.add_image(WorksheetImage(str(image_path)), "D4")
                workbook.save(path)
                workbook.close()
                if suffix == ".xlsm":
                    add_macro_package(path)
                with ZipFile(path) as archive:
                    images = {name: archive.read(name) for name in archive.namelist() if name.startswith("xl/media/")}
                self.assertEqual(len(images), 2)
                self.excel.load(path, MAPPING)
                for row in (2, 3):
                    self.excel.mark_executed(iter([row]))
                    with ZipFile(path) as archive:
                        self.assertEqual({name: archive.read(name) for name in images}, images)
                    if suffix == ".xlsm":
                        self.assert_macro_package_preserved(path)
                self.excel.close()
                saved = load_workbook(path)
                try:
                    self.assertEqual([saved.active.cell(row, 6).value for row in (2, 3)], ["已执行", "已执行"])
                    self.assertEqual(saved["其他工作表"]["A1"].value, "保留内容")
                    self.assertEqual(len(saved.active._images), 1)
                    self.assertEqual(len(saved["其他工作表"]._images), 1)
                    self.assertEqual(saved["其他工作表"]._images[0].anchor._from.col, 3)
                finally:
                    saved.close()

    def test_failed_write_or_replace_preserves_source_and_allows_retry(self) -> None:
        path = self.base / "failure.xlsm"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        add_macro_package(path)
        self.excel.load(path, MAPPING)
        original = path.read_bytes()

        def fail_mid_save(archive, entry, data):
            archive.fp.write(b"incomplete archive")
            raise OSError("save interrupted")

        for method, error in (("save", "save interrupted"), ("replace", "file locked")):
            with self.subTest(method=method):
                # Persistence now patches the ZIP worksheet instead of using
                # Workbook.save; inject the same partial-write fault there.
                failure = patch.object(ZipFile, "writestr", fail_mid_save) if method == "save" else patch.object(
                    Path, "replace", side_effect=PermissionError("file locked"),
                )
                with failure:
                    with self.assertRaisesRegex(OSError, error):
                        self.excel.mark_executed([2])
                self.assertEqual(path.read_bytes(), original)
                self.assertIsNone(self.excel.sheet["F2"].value)
                self.assertEqual(list(self.base.glob(".screenshot-word-*")), [])
        self.excel.mark_executed([2])
        self.assertEqual(self.excel.sheet["F2"].value, "已执行")
        self.assert_macro_package_preserved(path)

    def test_status_save_preserves_other_changes_already_on_disk(self) -> None:
        path = self.base / "latest.xlsx"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        self.excel.load(path, MAPPING)
        other = load_workbook(path)
        other.active["H8"] = "后来保存的备注"
        other.save(path)
        other.close()
        self.excel.mark_executed([2])
        saved = load_workbook(path)
        try:
            self.assertEqual(saved.active["H8"].value, "后来保存的备注")
            self.assertEqual(saved.active["F2"].value, "已执行")
        finally:
            saved.close()

    def test_all_sheets_execute_in_tab_order_with_independent_mappings_and_inheritance(self) -> None:
        path = self.base / "all_sheets.xlsm"
        workbook = Workbook()
        first = workbook.active
        first.title = "第一张"
        first.append(HEADERS)
        first.append(["登录", "成功", "步骤1", "", "", ""])
        first.append([None, None, "步骤2", "", "", ""])
        first.merge_cells("A2:A3")
        first.merge_cells("B2:B3")
        empty = workbook.create_sheet("空白表")
        empty["F20"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        second = workbook.create_sheet("第二张")
        second.append(["Status", "Check Point", "Case Name"])
        second.append([None, "独立验证点", "支付"])
        second.append(["已执行", "完成", "退出"])
        workbook.active = 2
        workbook.save(path)
        workbook.close()
        add_macro_package(path)

        previews = self.excel.preview_all(path)
        self.assertEqual([p.sheet_name for p in previews], ["第一张", "第二张"])
        # The legacy single-sheet entry points are deterministic as well.
        self.assertEqual(self.excel.preview(path).sheet_name, "第一张")
        mappings = {p.sheet_name: p.suggested_mapping for p in reversed(previews)}
        items = self.excel.load_all(path, mappings)
        self.assertEqual([(i.sheet_name, i.row_index, i.case_name) for i in items], [
            ("第一张", 2, "登录"), ("第一张", 3, "登录"), ("第二张", 2, "支付"),
        ])
        with self.assertRaisesRegex(ValueError, "必须指定工作表"):
            self.excel.mark_executed([2])
        for item in items:
            self.excel.mark_executed([item.row_index], sheet_name=item.sheet_name)
        saved = load_workbook(path, keep_vba=True)
        try:
            self.assertEqual(saved.active.title, "第二张")
            self.assertEqual([saved["第一张"].cell(row, 6).value for row in (2, 3)], ["已执行"] * 2)
            self.assertEqual(saved["第二张"]["A2"].value, "已执行")
            self.assertEqual(saved["第二张"]["C2"].value, "支付")
            self.assertIsNone(saved["空白表"]["F20"].value)
        finally:
            self.excel._close_workbook(saved)
        self.assert_macro_package_preserved(path)
        self.assertEqual(self.excel.load_all(path, mappings), [])

    def test_all_sheets_require_valid_mappings_and_release_workbook_on_failure(self) -> None:
        path = self.base / "invalid_sheet.xlsm"
        make_workbook(path, [["登录", "成功", "", "", "", ""]])
        workbook = load_workbook(path)
        workbook.create_sheet("第二张").append(["不规范标题", "需要手动映射"])
        workbook.save(path)
        workbook.close()
        add_macro_package(path)
        original = path.read_bytes()
        for mappings, error in (({"Sheet": MAPPING}, "工作表已变化"),
                                ({"Sheet": MAPPING, "第二张": {}}, "第二张")):
            with self.subTest(error=error):
                with self.assertRaisesRegex(ValueError, error):
                    self.excel.load_all(path, mappings)
                self.assertIsNone(self.excel.workbook)
                self.assertEqual(path.read_bytes(), original)

    def test_status_patch_does_not_reserialize_other_parts_or_reload_workbook(self) -> None:
        path = self.base / "parts.xlsm"
        make_workbook(path, [["登录", "成功", "步骤1", "", "", ""],
                             [None, None, "步骤2", "", "", ""]])
        add_macro_package(path)
        with ZipFile(path, "a") as archive:
            archive.writestr("customXml/preserved.bin", b"preserve-unknown-extension\x00\xff")
            archive.comment = b"package comment"
        with ZipFile(path) as archive:
            before = {name: archive.read(name) for name in archive.namelist()}
        self.excel.load(path, MAPPING)
        with patch("qt_app.services.excel_service.load_workbook", side_effect=AssertionError("unexpected reload")), \
             patch.object(Workbook, "save", side_effect=AssertionError("unexpected save")):
            self.excel.mark_executed([2])
            self.excel.mark_executed([3])
        with ZipFile(path) as archive:
            self.assertEqual(archive.comment, b"package comment")
            self.assertEqual(len(archive.namelist()), len(set(archive.namelist())))
            after = {name: archive.read(name) for name in archive.namelist()}
        self.assertEqual(set(before), set(after))
        self.assertEqual([name for name in before if before[name] != after[name]], ["xl/worksheets/sheet1.xml"])
        self.assert_macro_package_preserved(path)

    def test_status_patch_keeps_cell_styles_namespace_extensions_and_cell_order(self) -> None:
        path = self.base / "styled.xlsx"
        make_workbook(path, [["登录", "成功", "", "", "", "待执行", "保留右侧值"]])
        workbook = load_workbook(path)
        workbook.active["F2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        workbook.save(path)
        workbook.close()
        self.excel.load(path, MAPPING)
        self.excel.mark_executed([2], status=" 通过 & <确认> ")
        saved = load_workbook(path)
        try:
            self.assertEqual(saved.active["F2"].fill.fgColor.rgb, "00FFFF00")
            self.assertEqual(saved.active["F2"].value, " 通过 & <确认> ")
            self.assertEqual(saved.active["G2"].value, "保留右侧值")
        finally:
            saved.close()
        xml = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
               b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
               b'xmlns:x14="urn:test-extension" mc:Ignorable="x14">'
               b'<sheetData><row r="2"><c r="A2"/><c r="G2"/></row></sheetData>'
               b'<extLst><ext uri="extension"><x14:unchanged/></ext></extLst></worksheet>')
        patched, _ = self.excel._patch_status_xml(xml, [2], 6, "已执行")
        root = etree.fromstring(patched)
        self.assertEqual(root.nsmap["x14"], "urn:test-extension")
        self.assertIsNotNone(root.find("{*}extLst/{*}ext/{urn:test-extension}unchanged"))
        self.assertEqual([c.get("r") for c in root.find("{*}sheetData/{*}row")], ["A2", "F2", "G2"])

    def test_status_patch_handles_merged_results_and_missing_sheet_without_touching_source(self) -> None:
        path = self.base / "merged_result.xlsx"
        make_workbook(path, [["登录", "成功", "步骤1", "", "", ""],
                             [None, None, "步骤2", "", "", ""]])
        workbook = load_workbook(path)
        workbook.active.merge_cells("F2:F3")
        workbook.save(path)
        workbook.close()
        self.excel.load(path, MAPPING)
        self.excel.mark_executed([3])
        saved = load_workbook(path)
        try:
            self.assertEqual(saved.active["F2"].value, "已执行")
            self.assertIsNone(saved.active["F3"].value)
            saved.active.title = "已更名"
            saved.save(path)
        finally:
            saved.close()
        original = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "工作表已变化"):
            self.excel.mark_executed([2])
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(list(self.base.glob(".screenshot-word-*")), [])

    def test_external_save_during_status_write_is_preserved_and_retry_succeeds(self) -> None:
        path = self.base / "concurrent.xlsx"
        make_workbook(path, [["登录", "成功", "步骤1", "", "", ""]])
        self.excel.load(path, MAPPING)
        external_bytes = None

        def save_concurrently(reader, writer, length):
            nonlocal external_bytes
            copyfileobj(reader, writer, length)
            if external_bytes is None:
                workbook = load_workbook(path)
                workbook.active["H8"] = "外部新备注"
                replacement = self.base / "external.xlsx"
                workbook.save(replacement)
                workbook.close()
                replacement.replace(path)
                external_bytes = path.read_bytes()

        with patch("qt_app.services.excel_service.copyfileobj", side_effect=save_concurrently):
            with self.assertRaisesRegex(OSError, "其他程序修改"):
                self.excel.mark_executed([2])
        self.assertEqual(path.read_bytes(), external_bytes)
        self.assertIsNone(self.excel.sheet["F2"].value)
        self.assertEqual(list(self.base.glob(".screenshot-word-*")), [])
        self.excel.mark_executed([2])
        saved = load_workbook(path)
        try:
            self.assertEqual(saved.active["H8"].value, "外部新备注")
            self.assertEqual(saved.active["F2"].value, "已执行")
        finally:
            saved.close()

    def test_changed_merged_result_layout_fails_before_replacing_source(self) -> None:
        path = self.base / "changed_merge.xlsx"
        make_workbook(path, [["登录", "成功", "步骤1", "", "", ""],
                             [None, None, "步骤2", "", "", ""]])
        self.excel.load(path, MAPPING)
        workbook = load_workbook(path)
        workbook.active.merge_cells("F2:F3")
        workbook.save(path)
        workbook.close()
        original = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "合并单元格已变化"):
            self.excel.mark_executed([3])
        self.assertEqual(path.read_bytes(), original)
        self.assertIsNone(self.excel.sheet["F3"].value)
        self.assertEqual(list(self.base.glob(".screenshot-word-*")), [])

    def test_empty_and_whitespace_rows_are_skipped_before_previous_row_inheritance(self) -> None:
        path = self.base / "empty_rows.xlsx"
        make_workbook(path, [
            ["登录", "成功", "步骤1", "", "", ""],
            [None] * 6,
            ["", "", "步骤2", "", "", ""],
            [" \t", "\n", "\u3000", "\t", " ", "待执行"],
        ])
        workbook = load_workbook(path)
        workbook.active["A6"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        workbook.save(path)
        workbook.close()

        items = self.excel.load(path, MAPPING)
        self.assertEqual([item.row_index for item in items], [2, 4])
        self.assertEqual([item.step_name for item in items], ["步骤1", "步骤2"])
        self.assertTrue(all(item.case_name == "登录" and item.checkpoint == "成功" for item in items))

    def test_merged_blank_rows_are_skipped_but_real_steps_keep_inheritance(self) -> None:
        path = self.base / "merged_rows.xlsx"
        make_workbook(path, [
            ["登录", "成功", "前置", "", "", "已执行"],
            [None] * 6,
            [None, None, " \t", "\n", "\u3000", ""],
            [None, None, "步骤1", "", "", ""],
            ["退出", "", "步骤2", "", "", ""],
            [None, None, "", "真实步骤描述", "", ""],
            [None, None, "", "", "真实预期结果", ""],
        ])
        workbook = load_workbook(path)
        workbook.active.merge_cells("A2:A5")
        workbook.active.merge_cells("B2:B5")
        workbook.save(path)
        workbook.close()

        items = self.excel.load(path, MAPPING)
        self.assertEqual([item.row_index for item in items], [5, 6, 7, 8])
        self.assertEqual((items[0].case_name, items[0].checkpoint), ("登录", "成功"))
        self.assertTrue(all(item.case_name == "退出" for item in items[1:]))
        self.assertTrue(all(item.checkpoint == "未填写验证点" for item in items[1:]))

    def test_executed_status_is_checked_per_step_after_case_inheritance(self) -> None:
        path = self.base / "status_rows.xlsx"
        make_workbook(path, [
            ["登录", "成功", "前置", "", "", " PASS "],
            ["", "", "步骤1", "", "", ""],
            ["", "", "步骤2", "", "", "已执行"],
            ["", "", "步骤3", "", "", "失败"],
        ])
        items = self.excel.load(path, MAPPING)
        # PASS is a business outcome, so this row remains pending too.
        self.assertEqual([item.row_index for item in items], [2, 3, 5])
        self.assertTrue(all(item.case_name == "登录" and item.checkpoint == "成功" for item in items))

    def test_business_result_values_remain_pending(self) -> None:
        path = self.base / "business_results.xlsx"
        make_workbook(path, [
            ["登录", "登录成功", "步骤1", "", "", "成功"],
            ["支付", "支付完成", "步骤1", "", "", "通过"],
            ["退出", "退出成功", "步骤1", "", "", "pass"],
            ["查询", "查询结果", "步骤1", "", "", "passed"],
            ["注册", "注册成功", "步骤1", "", "", "已执行"],
        ])
        items = self.excel.load(path, MAPPING)
        self.assertEqual([item.row_index for item in items], [2, 3, 4, 5])


if __name__ == "__main__":
    unittest.main()
