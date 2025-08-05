import os
import time
import pandas as pd
import pyautogui
import keyboard
from tkinter import Tk, Canvas, PhotoImage, simpledialog, Button, Frame, Label, StringVar, Entry, messagebox
from PIL import Image, ImageTk
from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook
import glob
import re
import tempfile
import uuid

# 定义常量
INPUT_DIR = 'excel_input'
OUTPUT_DIR = 'output'
EXECUTED_STATUS = '已执行'

# 自动读取 excel_input 目录下最新的 Excel 文件
xlsx_files = sorted(glob.glob(os.path.join(INPUT_DIR, '*.xlsx')), key=os.path.getmtime, reverse=True)

if not xlsx_files:
    raise FileNotFoundError("在 ./excel_input/ 下未找到 Excel 文件，请放置一个案例文件")

EXCEL_PATH = xlsx_files[0]
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 加载Excel
df = pd.read_excel(EXCEL_PATH, engine='openpyxl')
df.fillna('', inplace=True)

# 标注工具类
class Annotator:
    def __init__(self, image_path, save_path, word_path, case_text):
        self.root = Tk()
        self.root.title("截图标注工具")
        self.image = Image.open(image_path)
        self.tk_image = ImageTk.PhotoImage(self.image)
        self.canvas = Canvas(self.root, width=self.image.width, height=self.image.height)
        self.canvas.pack()
        self.canvas.create_image(0, 0, anchor='nw', image=self.tk_image)

        self.start_x = None
        self.start_y = None
        self.drawn_objects = []
        self.text_mode = False
        self.save_path = save_path
        self.word_path = word_path
        self.case_text = case_text

        self.canvas.bind('<ButtonPress-1>', self.on_click)
        self.canvas.bind('<B1-Motion>', self.on_drag)
        self.canvas.bind('<ButtonRelease-1>', self.on_release)
        self.canvas.bind('<Button-3>', self.on_right_click)

        ctrl = Frame(self.root)
        ctrl.pack()
        Button(ctrl, text='撤销', command=self.undo).pack(side='left')
        Button(ctrl, text='保存截图并退出', command=self.save_and_exit).pack(side='left')

        self.root.mainloop()

    def on_click(self, event):
        if self.text_mode:
            return
        self.start_x = event.x
        self.start_y = event.y
        self.current_shape = self.canvas.create_oval(event.x, event.y, event.x, event.y, outline='red', width=2)
        self.drawn_objects.append(self.current_shape)

    def on_drag(self, event):
        if self.text_mode:
            return
        self.canvas.coords(self.current_shape, self.start_x, self.start_y, event.x, event.y)

    def on_release(self, event):
        pass

    def on_right_click(self, event):
        self.text_mode = True
        text = simpledialog.askstring("输入文字", "请输入标注内容：")
        if text:
            label = self.canvas.create_text(event.x, event.y, text=text, fill="blue", font=("Arial", 14, "bold"))
            self.drawn_objects.append(label)
        self.text_mode = False

    def undo(self):
        if self.drawn_objects:
            self.canvas.delete(self.drawn_objects.pop())

    def save_and_exit(self):
        temp_eps_fd, temp_eps_path = tempfile.mkstemp(suffix='.eps')
        try:
            self.canvas.postscript(file=temp_eps_path)
            img = Image.open(temp_eps_path)
            img.save(self.save_path)
        except Exception as e:
            messagebox.showerror("保存错误", f"保存截图时出错: {str(e)}")
            return
        finally:
            os.close(temp_eps_fd)
            if os.path.exists(temp_eps_path):
                os.remove(temp_eps_path)

        try:
            if os.path.exists(self.word_path):
                doc = Document(self.word_path)
            else:
                doc = Document()
                doc.add_heading(os.path.splitext(os.path.basename(self.word_path))[0], level=1)
                doc.add_paragraph(self.case_text)

            doc.add_picture(self.save_path, width=Inches(6))
            doc.add_paragraph('')
            doc.save(self.word_path)
        except Exception as e:
            messagebox.showerror("Word保存错误", f"保存Word文档时出错: {str(e)}")
            return

        self.root.destroy()

# 主处理流程
def run():
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active

    for idx, row in df.iterrows():
        if idx == 0:
            continue
        file_name = str(row[0]).strip()
        verify_point = str(row[1]).strip()
        status = str(row[2]).strip()

        # 路径安全检查，防止路径遍历攻击
        if re.search(r'[<>:"/\\|?*\x00-\x1F]', file_name):
            messagebox.showerror("文件名错误", f"无效的文件名: {file_name}")
            continue

        # 构造安全路径
        safe_file_name = os.path.normpath(file_name).replace(os.sep, '_')
        raw_dir = os.path.join(OUTPUT_DIR, safe_file_name)
        os.makedirs(raw_dir, exist_ok=True)

        if status.lower() in ['passed', '已执行']:
            continue

        # 显示验证点内容
        print(f"\n[用例 {file_name}] 验证点: {verify_point}")
        messagebox.showinfo("验证点", verify_point)

        while True:
            print("请按 F8 截图并标注...")
            keyboard.wait('F8')
            time.sleep(0.5)

            try:
                screenshot = pyautogui.screenshot()
            except Exception as e:
                messagebox.showerror("截图错误", f"截图时出错: {str(e)}")
                continue

            # 更高效的计数方式
            count = len(glob.glob(os.path.join(raw_dir, "screenshot_*.png"))) + 1
            img_path = os.path.join(raw_dir, f"screenshot_{count}.png")

            word_path = os.path.join(raw_dir, f"{safe_file_name}.docx")
            tmp_path = os.path.join(raw_dir, f"temp_{uuid.uuid4().hex}.png")
            screenshot.save(tmp_path)

            Annotator(tmp_path, img_path, word_path, verify_point)
            os.remove(tmp_path)

            ws.cell(row=idx+2, column=3, value=EXECUTED_STATUS)
            try:
                wb.save(EXCEL_PATH)
            except Exception as e:
                messagebox.showerror("保存Excel错误", f"保存Excel时出错: {str(e)}")
                continue

            choice = messagebox.askyesno("继续截图？", "是否继续截图？\n是 = 继续当前用例截图\n否 = 进入下一条验证点")
            if not choice:
                break

    messagebox.showinfo("执行完成", "所有用例执行完毕！")

if __name__ == '__main__':
    run()