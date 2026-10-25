from __future__ import annotations

import hashlib
import re
from tempfile import NamedTemporaryFile
import unicodedata
from pathlib import Path

INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]+')


class WordService:
    def __init__(self, output_dir: Path, logger):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.logger = logger

    def add_image(
        self, case_name: str, checkpoint: str, image_path: Path, step_note: str = "",
        *, sheet_name: str = "", operation_id: str = "",
    ) -> Path:
        # Lazy import: python-docx is only needed when the user actually finishes a screenshot.
        from docx import Document
        from docx.shared import Mm

        doc_path, doc = self._resolve_report(case_name, sheet_name=sheet_name)
        if doc is not None:
            self.logger.log(f"打开已有 Word 报告：{doc_path}")
        else:
            doc = Document()
            doc.add_heading(case_name, level=1)
            self.logger.log(f"创建 Word 报告：{doc_path}")
        identity = self._case_identifier(case_name) if not sheet_name else self._sheet_case_identifier(sheet_name, case_name)
        doc.core_properties.identifier = identity

        marker = f"screenshot-to-word-entry:{operation_id}" if operation_id else ""
        if marker and self._has_entry(doc, marker):
            return doc_path

        checkpoint_paragraph = self._find_checkpoint(doc, checkpoint) if checkpoint else None
        if checkpoint and checkpoint_paragraph is None:
            checkpoint_paragraph = doc.add_paragraph(f"验证点：{checkpoint}")
        section_end = self._next_checkpoint_element(checkpoint_paragraph) if checkpoint_paragraph is not None else None
        new_elements = []
        if step_note:
            for line in step_note.splitlines():
                if line.strip():
                    new_elements.append(doc.add_paragraph(line.strip())._p)
        picture_paragraph = doc.add_paragraph()
        inline_shape = picture_paragraph.add_run().add_picture(str(image_path), width=Mm(142.24))  # 5.6 in, within default A4 text width
        if marker:
            inline_shape._inline.docPr.set("descr", marker)
        new_elements.append(picture_paragraph._p)
        new_elements.append(doc.add_paragraph()._p)
        if section_end is not None:
            for element in new_elements:
                section_end.addprevious(element)
        temporary_path = None
        try:
            with NamedTemporaryFile(dir=self.output_dir, prefix=".screenshot-word-", suffix=".docx", delete=False) as temporary:
                temporary_path = Path(temporary.name)
            doc.save(str(temporary_path))
            temporary_path.replace(doc_path)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        self.logger.log(f"截图已写入 Word：{doc_path}")
        return doc_path

    def _resolve_report(self, case_name: str, *, sheet_name: str = ""):
        from docx import Document

        base = self._safe_filename(case_name)
        identity = self._case_identifier(case_name) if not sheet_name else self._sheet_case_identifier(sheet_name, case_name)
        digest = identity.rsplit(":", 1)[1][:12]

        def candidate(index: int) -> str:
            if index == 0:
                return f"{base}.docx"
            suffix = f"-{digest}" + (f"-{index}" if index > 1 else "")
            return f"{base[:80 - len(suffix)].rstrip(' ._')}{suffix}.docx"

        # Treat case/Unicode-equivalent names as occupied on every platform. This
        # also protects reports copied between case-sensitive and macOS/Windows disks.
        existing = {}
        for path in self.output_dir.iterdir():
            existing.setdefault(self._filename_key(path.name), []).append(path)

        indices = {0, 1}
        suffix_pattern = re.compile(rf"-{digest}(?:-(\d+))?\.docx$", re.IGNORECASE)
        for filename in existing:
            match = suffix_pattern.search(filename)
            if match:
                index = int(match.group(1) or "1")
                if index > 0 and filename == self._filename_key(candidate(index)):
                    indices.add(index)

        # Find our report before choosing a free name, including after another
        # colliding report was removed. The heading is the identity used by old reports.
        for index in sorted(indices):
            paths = existing.get(self._filename_key(candidate(index)), [])
            if len(paths) != 1 or not paths[0].is_file():
                continue
            try:
                doc = Document(str(paths[0]))
            except Exception:
                # An unreadable/foreign file must never be overwritten or appended to.
                continue
            stored_identity = doc.core_properties.identifier
            if stored_identity:
                matches_case = stored_identity == identity
            elif sheet_name:
                # Legacy reports have no sheet identity; do not guess ownership
                # and accidentally append a different sheet's case to them.
                matches_case = False
            else:
                heading = next(
                    (p.text for p in doc.paragraphs if p.style and p.style.style_id == "Heading1"),
                    None,
                )
                matches_case = heading == case_name
            if matches_case:
                return paths[0], doc

        index = 0
        while self._filename_key(candidate(index)) in existing:
            index += 1
        return self.output_dir / candidate(index), None

    @staticmethod
    def _filename_key(name: str) -> str:
        return unicodedata.normalize("NFC", name).casefold()

    @staticmethod
    def _case_identifier(case_name: str) -> str:
        digest = hashlib.sha256(case_name.encode("utf-8")).hexdigest()
        return f"screenshot-to-word:case-sha256:{digest}"

    @staticmethod
    def _sheet_case_identifier(sheet_name: str, case_name: str) -> str:
        digest = hashlib.sha256(f"{sheet_name}\0{case_name}".encode("utf-8")).hexdigest()
        return f"screenshot-to-word:sheet-case-sha256:{digest}"

    @staticmethod
    def _has_entry(doc, marker: str) -> bool:
        return any(shape._inline.docPr.get("descr") == marker for shape in doc.inline_shapes)

    @staticmethod
    def _find_checkpoint(doc, checkpoint: str):
        expected = f"验证点：{checkpoint}".strip()
        return next((paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == expected), None)

    @classmethod
    def _next_checkpoint_element(cls, paragraph):
        if paragraph is None:
            return None
        element = paragraph._p.getnext()
        while element is not None:
            text = "".join(element.xpath(".//w:t/text()"))
            if text.startswith("验证点："):
                return element
            element = element.getnext()
        return None

    def _safe_filename(self, name: str) -> str:
        cleaned = INVALID_FILENAME_CHARS.sub("_", name).strip(" ._")
        return cleaned[:80].rstrip(" ._") or "未命名用例"
