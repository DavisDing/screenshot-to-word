from __future__ import annotations

from dataclasses import dataclass
from copy import copy
from pathlib import Path
import os
import posixpath
from shutil import copymode, copyfileobj
from tempfile import NamedTemporaryFile
from typing import Dict, Iterable, List, Optional, Tuple
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from qt_app.models.test_case import TestItem
from utils.excel_schema import FIELD_SPECS, STEP_HINT_FIELDS, suggest_column_mapping

EXECUTED_STATUSES = {"已执行", "pass", "passed", "通过", "done", "completed", "完成", "成功"}
ROLE_CASE = "测试名称"
ROLE_CHECKPOINT = "验证点"
ROLE_STEP_NAME = "步骤名称"
ROLE_STEP_DESC = "步骤描述"
ROLE_EXPECTED = "预期结果"
ROLE_RESULT = "测试结果"
REQUIRED_ROLES = (ROLE_CASE, ROLE_CHECKPOINT, ROLE_RESULT)


@dataclass(frozen=True)
class ColumnInfo:
    index: int
    header: str
    samples: Tuple[str, ...]


@dataclass(frozen=True)
class ExcelPreview:
    file_path: Path
    sheet_name: str
    header_row: int
    columns: Tuple[ColumnInfo, ...]
    suggested_mapping: Dict[str, Optional[str]]
    reasons: Dict[str, str]

    @property
    def missing_required_roles(self) -> List[str]:
        return [role for role in REQUIRED_ROLES if not self.suggested_mapping.get(role)]


class ExcelService:
    """Excel reader/writer with tolerant column recognition and manual mapping support."""

    def __init__(self, logger):
        self.logger = logger
        self.file_path: Optional[Path] = None
        self.workbook = None
        self.sheet: Optional[Worksheet] = None
        self.header_row = 1
        self.columns: List[ColumnInfo] = []
        self.column_map: Dict[str, int] = {}
        self._merged_parent_cache: Dict[Tuple[int, int], Tuple[int, int]] = {}
        self.version = "基础版"
        self._sheet_mappings: Dict[str, Dict[str, Optional[str]]] = {}

    def preview(self, file_path: Path) -> ExcelPreview:
        file_path = Path(file_path)
        wb = load_workbook(file_path, keep_vba=file_path.suffix.lower() == ".xlsm")
        try:
            return self._preview_sheet(file_path, wb.worksheets[0])
        finally:
            self._close_workbook(wb)

    def preview_all(self, file_path: Path) -> List[ExcelPreview]:
        file_path = Path(file_path)
        wb = load_workbook(file_path, keep_vba=file_path.suffix.lower() == ".xlsm")
        try:
            return [self._preview_sheet(file_path, ws) for ws in wb.worksheets if self._has_content(ws)]
        finally:
            self._close_workbook(wb)

    def _preview_sheet(self, file_path: Path, ws: Worksheet) -> ExcelPreview:
        header_row = self._detect_header_row(ws)
        columns = self._read_columns(ws, header_row)
        suggestions, reasons = suggest_column_mapping([c.header for c in columns])
        return ExcelPreview(file_path, ws.title, header_row, tuple(columns), suggestions, reasons)

    @staticmethod
    def _has_content(ws: Worksheet) -> bool:
        return any(str(value).strip() for row in ws.iter_rows(values_only=True) for value in row if value is not None)

    def load(self, file_path: Path, mapping_by_role: Dict[str, Optional[str]]) -> List[TestItem]:
        """Compatibility entry point for a single table, always on the first worksheet."""
        return self._load(file_path, mapping_by_role=mapping_by_role)

    def load_all(self, file_path: Path, mappings_by_sheet: Dict[str, Dict[str, Optional[str]]]) -> List[TestItem]:
        return self._load(file_path, mappings_by_sheet=mappings_by_sheet)

    def _load(self, file_path: Path, *, mapping_by_role=None, mappings_by_sheet=None) -> List[TestItem]:
        self.close()
        self.file_path = Path(file_path)
        self.workbook = load_workbook(self.file_path, keep_vba=self.file_path.suffix.lower() == ".xlsm")
        try:
            if mappings_by_sheet is None:
                self._sheet_mappings = {self.workbook.worksheets[0].title: dict(mapping_by_role)}
            else:
                names = [ws.title for ws in self.workbook.worksheets if self._has_content(ws)]
                if set(names) != set(mappings_by_sheet):
                    raise ValueError("Excel 工作表已变化，请重新确认每张工作表的列映射")
                self._sheet_mappings = {name: dict(mappings_by_sheet[name]) for name in names}
            items = []
            for name in self._sheet_mappings:
                self._select_sheet(name)
                items.extend(self._iter_pending_items())
            if self._sheet_mappings:
                self._select_sheet(next(iter(self._sheet_mappings)))
        except Exception:
            self.close()
            raise
        self.logger.log(f"Excel 加载完成：{self.file_path}，工作表：{len(self._sheet_mappings)} 张，待执行：{len(items)} 条")
        return items

    def _select_sheet(self, name: str) -> None:
        if self.sheet is not None and self.sheet.title == name and self.column_map:
            return
        self.sheet = self.workbook[name]
        self._merged_parent_cache = {}
        self.header_row = self._detect_header_row(self.sheet)
        self.columns = self._read_columns(self.sheet, self.header_row)
        header_to_index = {c.header: c.index for c in self.columns}
        self.column_map = {}
        for role, header in self._sheet_mappings[name].items():
            if header and header in header_to_index:
                self.column_map[role] = header_to_index[header]

        missing = [role for role in REQUIRED_ROLES if role not in self.column_map]
        if missing:
            readable = "、".join(missing)
            raise ValueError(f"Excel 缺少必要列映射：{readable}（工作表：{name}）")

        self.version = "步骤版" if any(role in self.column_map for role in STEP_HINT_FIELDS) else "基础版"

    def mark_executed(self, row_indexes: Iterable[int], status: str = "已执行", *, sheet_name: Optional[str] = None) -> None:
        if not self.workbook or not self.sheet or not self.file_path:
            raise RuntimeError("Excel 尚未加载")
        rows = list(row_indexes)
        if not rows:
            return
        if any(not isinstance(row, int) or isinstance(row, bool) or row < 1 or row > 1048576 for row in rows):
            raise ValueError("Excel 行号无效")
        rows = sorted(set(rows))
        if sheet_name is None and len(self._sheet_mappings) > 1:
            raise ValueError("多工作表回写必须指定工作表名称")
        name = sheet_name if sheet_name is not None else self.sheet.title
        if name not in self._sheet_mappings:
            raise ValueError(f"Excel 工作表未加载：{name}")
        self._select_sheet(name)
        result_col = self.column_map[ROLE_RESULT]
        temporary_path = None
        try:
            with NamedTemporaryFile(
                dir=self.file_path.parent,
                prefix=".screenshot-word-",
                suffix=self.file_path.suffix,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            # XLSX/XLSM are ZIP packages: only the target worksheet XML is
            # edited. Other part contents stay unchanged; ZIP entries are still
            # streamed/recompressed into the atomic replacement file.
            with self.file_path.open("rb") as source_file, ZipFile(source_file) as source:
                source_stat = os.fstat(source_file.fileno())
                target = self._worksheet_part(source, name)
                xml, cells = self._patch_status_xml(source.read(target), rows, result_col, status)
                if cells != {self._merged_parent(row, result_col) for row in rows}:
                    raise ValueError("Excel 合并单元格已变化，请重新加载")
                with ZipFile(temporary_path, "w") as destination:
                    destination.comment = source.comment
                    for entry in source.infolist():
                        if entry.filename == target:
                            destination.writestr(copy(entry), xml)
                        else:
                            with source.open(entry) as reader, destination.open(copy(entry), "w") as writer:
                                copyfileobj(reader, writer, 1024 * 1024)
            current_stat = self.file_path.stat()
            if (source_stat.st_dev, source_stat.st_ino, source_stat.st_size, source_stat.st_mtime_ns) != (
                current_stat.st_dev, current_stat.st_ino, current_stat.st_size, current_stat.st_mtime_ns
            ):
                raise OSError("Excel 在保存期间被其他程序修改，请重试")
            copymode(self.file_path, temporary_path)
            temporary_path.replace(self.file_path)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        # Keep the in-memory execution state unchanged until persistence succeeds.
        for row, column in cells:
            self.sheet.cell(row=row, column=column).value = status
        self.logger.log(f"Excel 已标记执行结果：{self.file_path}，工作表：{name}")

    @staticmethod
    def _worksheet_part(archive: ZipFile, sheet_name: str) -> str:
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        workbook = etree.fromstring(archive.read("xl/workbook.xml"), parser)
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser)
        sheet = next((s for s in workbook.findall("{*}sheets/{*}sheet") if s.get("name") == sheet_name), None)
        if sheet is None:
            raise ValueError(f"Excel 工作表已变化：{sheet_name}，请重新加载")
        relationship_id = next((value for key, value in sheet.attrib.items() if key.endswith("}id")), None)
        relationship = next((r for r in relationships if r.get("Id") == relationship_id), None)
        if relationship is None or relationship.get("TargetMode") == "External":
            raise ValueError(f"Excel 工作表关系无效：{sheet_name}")
        target = relationship.get("Target", "")
        part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
        if not part.startswith("xl/") or part not in archive.namelist():
            raise ValueError(f"Excel 工作表路径无效：{sheet_name}")
        return part

    @staticmethod
    def _patch_status_xml(xml: bytes, rows: List[int], column: int, status: str):
        root = etree.fromstring(xml, etree.XMLParser(resolve_entities=False, no_network=True))
        ns = etree.QName(root).namespace
        def tag(name: str) -> str:
            return f"{{{ns}}}{name}"

        data = root.find(tag("sheetData"))
        if data is None:
            raise ValueError("Excel 工作表缺少数据区域")
        cells = {(row, column) for row in rows}
        for merged in root.findall(f"{tag('mergeCells')}/{tag('mergeCell')}"):
            min_col, min_row, max_col, max_row = range_boundaries(merged.get("ref"))
            for row, col in list(cells):
                if min_row <= row <= max_row and min_col <= col <= max_col:
                    cells.remove((row, col))
                    cells.add((min_row, min_col))
        row_elements = {int(row.get("r")): row for row in data}
        for row, col in sorted(cells):
            row_element = row_elements.get(row)
            if row_element is None:
                raise ValueError(f"Excel 第 {row} 行已变化，请重新加载")
            reference = f"{get_column_letter(col)}{row}"
            cell = next((c for c in row_element if c.get("r") == reference), None)
            if cell is None:
                cell = etree.Element(tag("c"), r=reference)
                following = next((c for c in row_element if coordinate_to_tuple(c.get("r"))[1] > col), None)
                if following is None:
                    row_element.append(cell)
                else:
                    following.addprevious(cell)
            for child in list(cell):
                if child.tag in {tag("f"), tag("v"), tag("is")}:
                    cell.remove(child)
            cell.set("t", "inlineStr")
            inline = etree.Element(tag("is"))
            text = etree.SubElement(inline, tag("t"))
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text.text = status
            cell.insert(0, inline)
        return etree.tostring(root, encoding="utf-8"), cells

    def close(self) -> None:
        if self.workbook:
            self._close_workbook(self.workbook)
        self.workbook = None
        self.sheet = None
        self._merged_parent_cache = {}
        self._sheet_mappings = {}

    @staticmethod
    def _close_workbook(workbook) -> None:
        try:
            workbook.close()
        finally:
            # openpyxl 的 Workbook.close() 不会关闭用于保留宏的内存归档。
            if workbook.vba_archive is not None:
                workbook.vba_archive.close()

    def _detect_header_row(self, ws: Worksheet) -> int:
        best_row = 1
        best_score = -1
        max_probe = min(ws.max_row, 15)
        aliases = []
        for spec in FIELD_SPECS.values():
            aliases.extend([spec.label, *spec.aliases])
        normalized_aliases = {str(a).strip().lower() for a in aliases if a}

        for row in range(1, max_probe + 1):
            values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
            non_empty = [str(v).strip() for v in values if v not in (None, "")]
            if len(non_empty) < 2:
                continue
            score = len(non_empty)
            lowered = {v.lower() for v in non_empty}
            score += len(lowered & normalized_aliases) * 5
            if score > best_score:
                best_score = score
                best_row = row
        return best_row

    def _read_columns(self, ws: Worksheet, header_row: int) -> List[ColumnInfo]:
        columns: List[ColumnInfo] = []
        used_headers = set()
        for col in range(1, ws.max_column + 1):
            raw_header = ws.cell(row=header_row, column=col).value
            header = str(raw_header).strip() if raw_header not in (None, "") else f"第{col}列"
            if header in used_headers:
                header = f"{header}({col})"
            used_headers.add(header)
            samples = []
            for row in range(header_row + 1, min(ws.max_row, header_row + 5) + 1):
                value = ws.cell(row=row, column=col).value
                if value not in (None, ""):
                    samples.append(str(value).strip())
            columns.append(ColumnInfo(index=col, header=header, samples=tuple(samples[:3])))
        return columns

    def _get_cell_text(self, row: int, role: str, *, resolve_merged: bool = False) -> str:
        if not self.sheet or role not in self.column_map:
            return ""
        column = self.column_map[role]
        value = self.sheet.cell(row=row, column=column).value
        if value in (None, "") and resolve_merged:
            parent_row, parent_col = self._merged_parent(row, column)
            if (parent_row, parent_col) != (row, column):
                value = self.sheet.cell(row=parent_row, column=parent_col).value
        return str(value).strip() if value not in (None, "") else ""

    def _merged_parent(self, row: int, column: int) -> Tuple[int, int]:
        if not self.sheet:
            return row, column
        key = (row, column)
        if key in self._merged_parent_cache:
            return self._merged_parent_cache[key]
        parent = key
        for merged_range in self.sheet.merged_cells.ranges:
            if (merged_range.min_row <= row <= merged_range.max_row) and (merged_range.min_col <= column <= merged_range.max_col):
                parent = (merged_range.min_row, merged_range.min_col)
                break
        self._merged_parent_cache[key] = parent
        return parent

    def _mark_cell_value(self, row: int, column: int, value: str) -> None:
        if not self.sheet:
            return
        parent_row, parent_col = self._merged_parent(row, column)
        self.sheet.cell(row=parent_row, column=parent_col).value = value

    def _iter_pending_items(self):
        if not self.sheet:
            return

        current_case = ""
        current_checkpoint = ""
        for row in range(self.header_row + 1, self.sheet.max_row + 1):
            row_case = self._get_cell_text(row, ROLE_CASE)
            row_checkpoint = self._get_cell_text(row, ROLE_CHECKPOINT)
            step_name = self._get_cell_text(row, ROLE_STEP_NAME)
            step_description = self._get_cell_text(row, ROLE_STEP_DESC)
            expected_result = self._get_cell_text(row, ROLE_EXPECTED)
            # 必须先检查当前行自身的内容，避免空行从上一行或合并单元格继承后成为待执行项。
            if not any((row_case, row_checkpoint, step_name, step_description, expected_result)):
                continue

            raw_case = self._get_cell_text(row, ROLE_CASE, resolve_merged=True)
            raw_checkpoint = self._get_cell_text(row, ROLE_CHECKPOINT, resolve_merged=True)

            if raw_case and raw_case != current_case:
                current_case = raw_case
                # 新用例开始但验证点为空时，不沿用上一条用例的验证点。
                current_checkpoint = raw_checkpoint or ""
            elif raw_case:
                current_case = raw_case
            if raw_checkpoint:
                current_checkpoint = raw_checkpoint

            case_name = current_case
            checkpoint = current_checkpoint

            # 测试结果列按当前行判断；测试名称/验证点可从合并单元格或上一行继承。
            result = self._get_cell_text(row, ROLE_RESULT, resolve_merged=False).lower()
            if result in EXECUTED_STATUSES:
                continue
            yield TestItem(
                row_index=row,
                case_name=case_name or f"未命名用例_{row}",
                checkpoint=checkpoint or "未填写验证点",
                step_name=step_name,
                step_description=step_description,
                expected_result=expected_result,
                sheet_name=self.sheet.title,
            )
